"""
KAP Survey Analyzer
Knowledge-Attitude-Practice survey analysis: scoring, Bloom's cut-offs, reliability,
associations, logistic regression, and a publication-ready Word/Excel report.
"""
import io
import re

import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

import kap_core as kc
import reporting as rp
from demo_data import ATTITUDE_REVERSE, CODEBOOK, KNOWLEDGE_KEY, make_demo

# ---- edit these for your branding -------------------------------------------
APP_NAME = "KAP Survey Analyzer"
AUTHOR = "Tooba Mujtaba"
BRAND = "NovaGen Bioinformatics"
CONTACT_URL = "https://www.linkedin.com/in/tooba-mujtabaa247ba190"
# -----------------------------------------------------------------------------

st.set_page_config(page_title=APP_NAME, page_icon="📋", layout="wide")


def show_df(d: pd.DataFrame):
    """Full-width table on both older and newer Streamlit versions."""
    try:
        st.dataframe(d, width="stretch", hide_index=True)
    except Exception:
        st.dataframe(d, use_container_width=True, hide_index=True)


def show_fig(fig):
    st.pyplot(fig)
    plt.close(fig)


DOMAINS = [("Knowledge", r"^(k|know\w*)[\s_.\-]*\d+"),
           ("Attitude", r"^(a|att\w*)[\s_.\-]*\d+"),
           ("Practice", r"^(p|prac\w*)[\s_.\-]*\d+")]
CUT_OPTIONS = {
    "bloom": "Bloom's cut-off (≥80% / 60–79% / <60%)",
    "two": "Custom three levels (two cut-offs)",
    "one": "Two levels at a fixed cut-off",
    "median": "Two levels at the median score",
    "mean": "Two levels at the mean score",
}


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------
def read_upload(file, sheet=None) -> pd.DataFrame:
    name = file.name.lower()
    data = file.getvalue()
    if name.endswith((".xlsx", ".xlsm")):
        df = pd.read_excel(io.BytesIO(data), sheet_name=sheet or 0)
    else:
        df = None
        for enc in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
            try:
                df = pd.read_csv(io.BytesIO(data), encoding=enc, sep=None, engine="python")
                break
            except (UnicodeDecodeError, pd.errors.ParserError):
                continue
        if df is None:
            raise ValueError("Could not read this CSV file. Save it as UTF-8 CSV or .xlsx and try again.")
    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[:, [c for c in df.columns if not c.startswith("Unnamed:") or df[c].notna().any()]]
    df = df.replace(r"^\s*$", np.nan, regex=True).dropna(how="all")
    return df.reset_index(drop=True)


def excel_sheets(file) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file.getvalue()), read_only=True)
    return wb.sheetnames


with st.sidebar:
    st.title(f"📋 {APP_NAME}")
    st.caption("Knowledge · Attitude · Practice survey analysis, from raw responses to publication tables.")
    up = st.file_uploader("Upload survey data (CSV or Excel)", type=["csv", "xlsx", "xlsm", "txt"])
    sheet = None
    if up is not None and up.name.lower().endswith((".xlsx", ".xlsm")):
        try:
            sheets = excel_sheets(up)
            if len(sheets) > 1:
                sheet = st.selectbox("Sheet", sheets)
        except Exception:
            pass
    if st.button("Try the demo dataset"):
        st.session_state["use_demo"] = True
    if up is not None:
        st.session_state["use_demo"] = False
    st.download_button("Download demo data (.csv)", make_demo().to_csv(index=False).encode(),
                       "demo_kap_survey.csv", "text/csv")
    st.divider()
    st.markdown("**Data format:** one row per respondent, one column per question. "
                "Keep the original answers (e.g. *True / False / Don't know*, *Agree*, *Always*); "
                "the app scores them for you.")
    st.caption("Your data is processed in this session only and is not stored by the app.")
    st.divider()
    st.markdown(f"Need the full analysis, manuscript Results section, or a custom questionnaire analysed? "
                f"[Contact {AUTHOR}]({CONTACT_URL})")
    st.caption(f"© {BRAND}")

df = None
source = ""
if up is not None:
    try:
        df = read_upload(up, sheet)
        source = up.name
    except Exception as e:
        st.error(f"Could not read the file: {e}")
elif st.session_state.get("use_demo"):
    df = make_demo()
    source = "demo dataset (simulated AMR KAP survey of 240 healthcare workers)"

if df is None:
    st.header(APP_NAME)
    st.markdown(
        "Analyse a **Knowledge, Attitude and Practice (KAP)** questionnaire in minutes:\n\n"
        "1. **Upload** your raw survey export (Google Forms, KoboToolbox, SurveyMonkey, Excel).\n"
        "2. **Map** the knowledge, attitude and practice questions and set the answer key or Likert coding.\n"
        "3. **Get** scores, Bloom's levels, KR-20/Cronbach's α, chi-square/Fisher tests, group comparisons, "
        "Spearman correlations and logistic regression (crude and adjusted ORs).\n"
        "4. **Download** a Word report with formatted tables, figures and an auto-written Methods and Results "
        "draft, plus an Excel workbook of every table.")
    st.info("Upload a file in the sidebar, or click **Try the demo dataset** to see a full example.")
    with st.expander("Example of the expected layout"):
        show_df(make_demo().head(6))
    st.stop()

N = len(df)
cols = list(df.columns)
st.subheader(f"{N} respondents · {len(cols)} columns")
st.caption(f"Source: {source}")

tab_vars, tab_score, tab_res, tab_rep = st.tabs(
    ["① Variables", "② Scoring", "③ Results", "④ Report & downloads"])

# ---------------------------------------------------------------------------
# ① Variables
# ---------------------------------------------------------------------------
is_demo = source.startswith("demo")
ss = st.session_state
data_sig = f"{source}|{N}|{len(cols)}"   # widget keys depend on the dataset, so a new file resets defaults


def guess_items(pattern: str) -> list[str]:
    return [c for c in cols if re.match(pattern, c.strip(), flags=re.I)]


with tab_vars:
    show_df(df.head(8))
    st.markdown("#### Questionnaire domains")
    st.caption("Pick the question columns for each domain. Columns named K1, A1, P1 (or Knowledge_1 …) are "
               "pre-selected. Rename a domain if your questionnaire uses e.g. *Awareness* or *Perception*.")
    domain_setup = []
    taken: set[str] = set()
    dcols = st.columns(3)
    for (dname, pat), dc in zip(DOMAINS, dcols):
        with dc:
            on = st.checkbox(f"Include {dname.lower()}", value=True, key=f"on_{dname}_{data_sig}")
            name = st.text_input("Domain name", value=dname, key=f"name_{dname}_{data_sig}", disabled=not on)
            items = st.multiselect("Questions", cols, default=guess_items(pat), key=f"items_{dname}_{data_sig}",
                                   disabled=not on)
            if on and items:
                domain_setup.append((dname, name.strip() or dname, items))
                taken.update(items)
    dup = [c for c in taken if sum(c in d[2] for d in domain_setup) > 1]
    if dup:
        st.warning(f"These columns are in more than one domain: {', '.join(dup)}")

    st.markdown("#### Participant characteristics")
    id_like = [c for c in cols if c not in taken and (re.fullmatch(r"(?i).*\bid\b.*|respondent|timestamp|.*time.*", c)
                                                       or (df[c].nunique() == N and not kc.is_continuous(df[c])))]
    cand = [c for c in cols if c not in taken and c not in id_like]
    chars = st.multiselect("Characteristics to describe and test", [c for c in cols if c not in taken],
                           default=cand, key=f"chars_{data_sig}")
    cont_default = [c for c in chars if kc.is_continuous(df[c])]
    continuous = st.multiselect("Treat as continuous (numeric)", chars, default=cont_default,
                                key=f"cont_{data_sig}",
                                help="Continuous variables are summarised as mean ± SD / median (IQR), correlated "
                                     "with scores (Spearman) and grouped for the cross-tabulations.")
    bins: dict[str, pd.Series] = {}
    if continuous:
        st.caption("Grouping of continuous variables for cross-tabulations and group comparisons:")
        bcols = st.columns(min(3, len(continuous)))
        for i, v in enumerate(continuous):
            with bcols[i % len(bcols)]:
                how = st.selectbox(v, ["Median split", "Custom cut points", "Do not group"],
                                   key=f"bin_{v}_{data_sig}")
                if how == "Custom cut points":
                    txt = st.text_input(f"Cut points for {v}", value="", placeholder="e.g. 30, 40",
                                        key=f"edges_{v}_{data_sig}")
                    edges = [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?", txt)]
                    if edges:
                        bins[v] = kc.bin_continuous(df[v], "custom", edges)
                    else:
                        st.caption("Enter at least one number; using the median until then.")
                        bins[v] = kc.bin_continuous(df[v], "median")
                elif how == "Median split":
                    bins[v] = kc.bin_continuous(df[v], "median")

if not domain_setup:
    with tab_score:
        st.info("Select the questions for at least one domain in **① Variables**.")
    st.stop()

# ---------------------------------------------------------------------------
# ② Scoring
# ---------------------------------------------------------------------------
configs: list[kc.DomainConfig] = []
with tab_score:
    st.caption("Set how each domain is scored and classified. Every choice is written into the Methods text.")
    for base, name, items in domain_setup:
        k = f"{base}_{data_sig}"
        uniq = sorted({str(v).strip() for it in items for v in df[it].dropna().unique()}, key=str.lower)
        normed = {kc.norm(u) for u in uniq}
        looks_key = normed <= {"true", "false", "yes", "no", "don't know", "dont know", "do not know",
                               "not sure", "i don't know", "correct", "incorrect", "0", "1", "maybe"}
        with st.expander(f"**{name}** · {len(items)} questions", expanded=True):
            method = st.radio("Scoring method", ["Answer key (correct = 1, otherwise 0)",
                                                 "Likert / rating scale"],
                              index=0 if looks_key else 1, horizontal=True, key=f"method_{k}")
            cfg = kc.DomainConfig(name=name, items=items, labels=kc.level_labels_for(base))
            if method.startswith("Answer"):
                cfg.method = "key"
                guess = kc.guess_correct_answer(uniq)
                opts = uniq
                g_index = opts.index(guess) if guess in opts else 0
                default_all = st.selectbox("Default correct answer (applied to every question below)", opts,
                                           index=g_index, key=f"defkey_{k}")
                st.caption("Change any question whose correct answer differs (e.g. false statements).")
                gcols = st.columns(4)
                for i, it in enumerate(items):
                    it_opts = sorted({str(v).strip() for v in df[it].dropna().unique()}, key=str.lower)
                    if default_all not in it_opts:
                        it_opts = [default_all] + it_opts
                    pre = default_all
                    if is_demo and it in KNOWLEDGE_KEY:
                        pre = KNOWLEDGE_KEY[it]
                    with gcols[i % 4]:
                        cfg.key[it] = st.selectbox(it, it_opts, index=it_opts.index(pre) if pre in it_opts else 0,
                                                   key=f"key_{k}_{it}_{default_all}")
            else:
                cfg.method = "likert"
                guess = kc.guess_mapping(uniq)
                map_df = pd.DataFrame({"Response": list(guess.keys()),
                                       "Score": [np.nan if v is None else v for v in guess.values()]})
                map_df = map_df.sort_values("Score", na_position="last").reset_index(drop=True)
                c1, c2 = st.columns([3, 2])
                with c1:
                    st.caption("Numeric value for each response (edit if needed; blank = treated as missing).")
                    edited = st.data_editor(map_df, hide_index=True,
                                            disabled=["Response"], key=f"map_{k}",
                                            column_config={"Score": st.column_config.NumberColumn(step=1)})
                    cfg.mapping = {r.Response: (None if pd.isna(r.Score) else float(r.Score))
                                   for r in edited.itertuples()}
                vals = [v for v in cfg.mapping.values() if v is not None]
                with c2:
                    lo = float(min(vals)) if vals else 1.0
                    hi = float(max(vals)) if vals else 5.0
                    cfg.scale_min = st.number_input("Scale minimum", value=lo, step=1.0, key=f"min_{k}")
                    cfg.scale_max = st.number_input("Scale maximum", value=hi, step=1.0, key=f"max_{k}")
                    rev_default = [i for i in ATTITUDE_REVERSE if i in items] if is_demo else []
                    cfg.reverse = st.multiselect("Reverse-scored (negatively worded) questions", items,
                                                 default=rev_default, key=f"rev_{k}")
                    cfg.pct_from_min = st.checkbox("Percent from the scale minimum ((score − min)/(max − min))",
                                                   value=False, key=f"pmin_{k}",
                                                   help="Default is score ÷ maximum possible × 100, the usual "
                                                        "convention with Bloom's cut-offs.")
                if cfg.scale_max <= cfg.scale_min:
                    st.error("Scale maximum must be greater than the minimum.")
                elif not cfg.pct_from_min and cfg.scale_min > 0:
                    floor = cfg.scale_min / cfg.scale_max * 100
                    st.caption(f"Note: with score ÷ maximum, the lowest possible percentage is {floor:.0f}%. "
                               "Tick *Percent from the scale minimum* if you want the full 0–100% range.")
            c1, c2 = st.columns([3, 2])
            with c1:
                cm = st.selectbox("Classification", list(CUT_OPTIONS), format_func=CUT_OPTIONS.get,
                                  key=f"cut_{k}")
                cfg.cutoff_method = cm
                if cm == "two":
                    a, b = st.columns(2)
                    cfg.cut_high = a.number_input("Upper cut-off (%)", 0.0, 100.0, 75.0, key=f"ch_{k}")
                    cfg.cut_low = b.number_input("Lower cut-off (%)", 0.0, 100.0, 50.0, key=f"cl_{k}")
                    if cfg.cut_low >= cfg.cut_high:
                        st.error("The lower cut-off must be below the upper cut-off.")
                elif cm == "one":
                    cfg.cut_high = st.number_input("Cut-off (%)", 0.0, 100.0, 60.0, key=f"c1_{k}")
            with c2:
                hi_l, mid_l, lo_l = cfg.labels
                l1, l2, l3 = st.columns(3)
                cfg.labels = (l1.text_input("High label", hi_l, key=f"lh_{k}"),
                              l2.text_input("Middle label", mid_l, key=f"lm_{k}",
                                            disabled=cm not in ("bloom", "two")),
                              l3.text_input("Low label", lo_l, key=f"ll_{k}"))
            configs.append(cfg)

    st.markdown("#### Analysis options")
    o1, o2 = st.columns(2)
    group_method = o1.radio("Group comparisons of % scores",
                            ["auto", "nonparametric", "parametric"],
                            format_func={"auto": "Automatic (Shapiro–Wilk decides)",
                                         "nonparametric": "Non-parametric (Mann–Whitney / Kruskal–Wallis)",
                                         "parametric": "Parametric (t-test / ANOVA)"}.get, key="gm")
    alpha_level = o2.number_input("Significance level (α)", 0.001, 0.2, 0.05, step=0.01, format="%.3f")

# score everything
results = [kc.score_domain(df, c) for c in configs]
for r in results:
    if r.unmapped:
        with tab_score:
            st.warning(f"**{r.cfg.name}**: some responses have no numeric value and were treated as missing - "
                       + "; ".join(f"{i}: {', '.join(v)}" for i, v in r.unmapped.items())
                       + ". Add a score for them in the table above.")
valid = [r for r in results if r.n_scored > 0]
if not valid:
    with tab_res:
        st.error("No respondent could be scored. Check the question selection and scoring settings.")
    st.stop()

# ---------------------------------------------------------------------------
# ③ Results
# ---------------------------------------------------------------------------
demo_tab = kc.demographics_table(df, chars, continuous) if chars else pd.DataFrame()
summary = kc.domain_summary(valid)
rel = kc.reliability_table(valid)
corr = kc.domain_correlations(valid)
assoc: dict[str, tuple[pd.DataFrame, list]] = {}
for r in valid:
    assoc[r.cfg.name] = kc.association_table(df, r, chars, continuous, bins) if chars else (pd.DataFrame(), [])
score_tab, score_tests = (kc.score_by_group_table(df, valid, chars, continuous, group_method, bins)
                          if chars else (pd.DataFrame(), []))
cont_corr = kc.continuous_corr_table(score_tests)
item_labels = dict(zip(CODEBOOK["Item"], CODEBOOK["Statement"])) if is_demo else {}

logit = None
with tab_res:
    rt = st.tabs(["Participants", "Items", "Scores & levels", "Reliability", "Associations",
                  "Score comparisons", "Correlations", "Logistic regression"])
    with rt[0]:
        if demo_tab.empty:
            st.info("Select participant characteristics in **① Variables**.")
        else:
            show_df(demo_tab)
    with rt[1]:
        for r in valid:
            st.markdown(f"##### {r.cfg.name}")
            show_df(kc.item_table(df, r).drop(columns=["% correct"], errors="ignore"))
            fig = (rp.fig_key_items(r, item_labels) if r.cfg.method == "key" else rp.fig_likert(df, r))
            if fig is not None:
                show_fig(fig)
    with rt[2]:
        mcols = st.columns(len(valid))
        for c, r in zip(mcols, valid):
            top = r.cfg.level_order[0]
            share = (r.level == top).sum() / r.n_scored * 100
            c.metric(f"{r.cfg.name}: {top.lower()}", f"{share:.1f}%",
                     help=f"Mean score {r.pct.mean():.1f}% · n = {r.n_scored}")
        show_df(summary)
        show_fig(rp.fig_levels(valid))
        excl = [f"{r.cfg.name}: {r.n_excluded}" for r in valid if r.n_excluded]
        if excl:
            st.caption("Respondents with any unanswered question in a domain were excluded from that domain "
                       f"({', '.join(excl)}).")
    with rt[3]:
        show_df(rel)
        st.caption("KR-20 is Cronbach's α for 0/1 items. ≥0.70 is generally acceptable; knowledge tests "
                   "with varied item difficulty often sit lower.")
        for r in valid:
            a = kc.cronbach_alpha(r.item_scores)
            if not a["items"].empty:
                with st.expander(f"Item statistics: {r.cfg.name}"):
                    show_df(a["items"])
                    low = a["items"].loc[a["items"]["Corrected item–total r"] < 0.2, "Item"].tolist()
                    if low:
                        st.caption(f"Items with corrected item–total r < 0.20: {', '.join(low)}. Check wording "
                                   "or answer key / reverse scoring.")
    with rt[4]:
        if not chars:
            st.info("Select participant characteristics in **① Variables**.")
        for r in valid:
            tab_a, tests = assoc[r.cfg.name]
            if tab_a.empty:
                continue
            st.markdown(f"##### {r.cfg.name} level by characteristic")
            show_df(tab_a)
            note = rp.test_note([tab_a])
            if note:
                st.caption(note + " Row percentages.")
    with rt[5]:
        if score_tab.empty:
            st.info("Select participant characteristics in **① Variables**.")
        else:
            show_df(score_tab)
            st.caption((rp.test_note([score_tab]) + " ") +
                       "Values are mean ± SD for parametric tests and median (IQR) otherwise, on the % score.")
            sig = [t for t in score_tests if t["test"] != "Spearman ρ" and pd.notna(t["p"]) and t["p"] < alpha_level]
            for t in sig:
                st.caption(f"{t['domain']} differs by {t['variable']} ({t['test']}, {kc.p_text(t['p'])}, "
                           f"{t['effect_name']} = {t['effect']:.2f}).")
    with rt[6]:
        if len(corr):
            show_df(corr.drop(columns=["_p", "_rho"]))
        else:
            st.info("Correlations need at least two domains.")
        if len(cont_corr):
            st.markdown("##### Continuous characteristics vs % scores")
            show_df(cont_corr)
    with rt[7]:
        preds_all = chars
        if not preds_all:
            st.info("Select participant characteristics in **① Variables**.")
        else:
            c1, c2 = st.columns(2)
            out_name = c1.selectbox("Outcome domain", [r.cfg.name for r in valid], key="lr_out")
            r_out = next(r for r in valid if r.cfg.name == out_name)
            pos = c2.multiselect("Outcome = 1 when level is", r_out.cfg.level_order,
                                 default=r_out.cfg.level_order[:1], key=f"lr_pos_{out_name}")
            preds = st.multiselect("Predictors", preds_all, default=preds_all, key=f"lr_preds_{data_sig}")
            refs = {}
            cat_preds = [p for p in preds if p not in continuous]
            if cat_preds:
                st.caption("Reference category for each categorical predictor:")
                rcols = st.columns(min(4, len(cat_preds)))
                for i, p in enumerate(cat_preds):
                    levels = kc._natural_order(df[p].dropna().astype(str).unique())
                    mode = df[p].dropna().astype(str).value_counts().idxmax() if levels else None
                    refs[p] = rcols[i % len(rcols)].selectbox(p, levels,
                                                              index=levels.index(mode) if mode in levels else 0,
                                                              key=f"ref_{p}_{data_sig}")
            if pos and len(pos) < len(r_out.cfg.level_order) and preds:
                y = r_out.level.isin(pos).astype(float).where(r_out.level.notna())
                label = " or ".join(p.lower() for p in pos) + f" {out_name.lower()}"
                try:
                    lt = kc.logistic_table(df, y, preds, continuous, refs)
                    lt["outcome_label"] = label
                    logit = lt
                    f = lt["fit"]
                    if f is not None:
                        m1, m2, m3, m4 = st.columns(4)
                        m1.metric("n in model", f.n)
                        m2.metric("Outcome events", f.events)
                        m3.metric("Model χ² p-value", kc.fmt_p(f.lr_p))
                        m4.metric("Nagelkerke R²", f"{f.nagelkerke:.3f}")
                    show_df(lt["table"].drop(columns=[c for c in lt["table"] if c.startswith("_")]))
                    st.caption(f"Outcome: {label} (1) vs other levels (0). COR = crude odds ratio from "
                               "univariable models; AOR = adjusted odds ratio from the multivariable model "
                               "including all listed predictors.")
                    for w in lt["warnings"]:
                        st.warning(w)
                except ValueError as e:
                    st.error(str(e))
            else:
                st.info("Choose at least one (but not every) outcome level and one predictor.")

# ---------------------------------------------------------------------------
# ④ Report
# ---------------------------------------------------------------------------
methods = kc.methods_text(valid, group_method, alpha_level, logit)
assoc_tests = {k: v[1] for k, v in assoc.items()}
results_paras = kc.results_text(N, valid, assoc_tests, corr, logit, alpha_level)

with tab_rep:
    title = st.text_input("Report title", value="Knowledge, attitude and practice survey: statistical report"
                          if not is_demo else "Knowledge, attitude and practice regarding antimicrobial "
                                              "resistance among healthcare workers (demo)")
    st.markdown("##### Statistical analysis (Methods draft)")
    st.text_area("methods", methods, height=220, label_visibility="collapsed")
    st.markdown("##### Results draft")
    st.text_area("results", "\n\n".join(results_paras), height=260, label_visibility="collapsed")
    st.caption("Drafts are generated from your settings and results. Review the wording against your study "
               "design before submission.")

    if st.button("Build Word report and Excel workbook", type="primary"):
        with st.spinner("Building report…"):
            blocks: list[dict] = []
            if not demo_tab.empty:
                blocks += [{"type": "heading", "text": "Participant characteristics"},
                           {"type": "table", "df": demo_tab,
                            "caption": f"Characteristics of the participants (N = {N})"}]
            for r in valid:
                blocks.append({"type": "heading", "text": r.cfg.name})
                it = kc.item_table(df, r).drop(columns=["% correct"], errors="ignore")
                blocks.append({"type": "table", "df": it,
                               "caption": f"{r.cfg.name} item responses",
                               "note": "(R) reverse-scored item." if r.cfg.reverse else ""})
                fig = rp.fig_key_items(r, item_labels) if r.cfg.method == "key" else rp.fig_likert(df, r)
                if fig is not None:
                    blocks.append({"type": "figure", "png": rp.fig_png(fig),
                                   "caption": f"{r.cfg.name} responses by item."})
            blocks.append({"type": "heading", "text": "KAP scores and levels"})
            for p in results_paras:
                if any(p.startswith(r.cfg.name + ":") for r in valid):
                    blocks.append({"type": "para", "text": p})
            blocks += [{"type": "table", "df": summary, "caption": "KAP scores and levels"},
                       {"type": "figure", "png": rp.fig_png(rp.fig_levels(valid)),
                        "caption": "Distribution of KAP levels by domain."},
                       {"type": "table", "df": rel, "caption": "Internal consistency of the questionnaire "
                                                               "domains"}]
            if chars:
                blocks.append({"type": "heading", "text": "Factors associated with KAP levels"})
                for p in results_paras:
                    if "significantly associated" in p:
                        blocks.append({"type": "para", "text": p})
                for r in valid:
                    tab_a = assoc[r.cfg.name][0]
                    blocks.append({"type": "table", "df": tab_a,
                                   "caption": f"Association between participant characteristics and "
                                              f"{r.cfg.name.lower()} level",
                                   "note": (rp.test_note([tab_a]) + " Row percentages.").strip()})
                blocks.append({"type": "table", "df": score_tab,
                               "caption": "Percentage scores by participant characteristics",
                               "note": (rp.test_note([score_tab]) + " Mean ± SD for parametric tests, "
                                        "median (IQR) otherwise.").strip()})
            if len(corr) or len(cont_corr):
                blocks.append({"type": "heading", "text": "Correlations"})
                for p in results_paras:
                    if p.startswith("Pairwise Spearman"):
                        blocks.append({"type": "para", "text": p})
                if len(corr):
                    blocks.append({"type": "table", "df": corr, "caption": "Spearman correlations between "
                                                                            "domain scores"})
                if len(cont_corr):
                    blocks.append({"type": "table", "df": cont_corr,
                                   "caption": "Spearman correlations of continuous characteristics with "
                                              "percentage scores"})
            if logit and logit.get("fit") is not None:
                blocks.append({"type": "heading", "text": "Logistic regression"})
                for p in results_paras:
                    if p.startswith("In multivariable"):
                        blocks.append({"type": "para", "text": p})
                f = logit["fit"]
                blocks.append({"type": "table", "df": logit["table"],
                               "caption": f"Factors associated with {logit['outcome_label']} "
                                          "(binary logistic regression)",
                               "note": f"COR, crude odds ratio; AOR, adjusted odds ratio; CI, confidence interval. "
                                       f"Model n = {f.n}, events = {f.events}; likelihood-ratio χ²({f.lr_df}) = "
                                       f"{f.lr_chi2:.2f}, {kc.p_text(f.lr_p)}; Nagelkerke R² = "
                                       f"{f.nagelkerke:.3f}."})
            ss["docx"] = rp.build_docx(title, N, methods, blocks)

            sections = [("Characteristics", demo_tab)]
            for r in valid:
                sections.append((f"{r.cfg.name} items", kc.item_table(df, r)))
            sections += [("Scores and levels", summary), ("Reliability", rel)]
            for r in valid:
                a = kc.cronbach_alpha(r.item_scores)["items"]
                sections.append((f"{r.cfg.name} item stats", a))
            for r in valid:
                sections.append((f"{r.cfg.name} associations", assoc[r.cfg.name][0]))
            sections += [("Scores by group", score_tab), ("Domain correlations", corr),
                         ("Continuous correlations", cont_corr)]
            if logit:
                sections.append(("Logistic regression", logit["table"]))
            scored = df.copy()
            for r in valid:
                scored[f"{r.cfg.name}_score"] = r.total
                scored[f"{r.cfg.name}_pct"] = r.pct.round(2)
                scored[f"{r.cfg.name}_level"] = r.level.astype(object)
            ss["xlsx"] = rp.build_excel(sections, methods, results_paras, scored)
            ss["scored_csv"] = scored.to_csv(index=False).encode()
            ss["built_for"] = data_sig
    if ss.get("docx") and ss.get("built_for") == data_sig:
        d1, d2, d3 = st.columns(3)
        d1.download_button("⬇ Word report (.docx)", ss["docx"], "KAP_report.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        d2.download_button("⬇ All tables (.xlsx)", ss["xlsx"], "KAP_tables.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        d3.download_button("⬇ Scored data (.csv)", ss["scored_csv"], "KAP_scored_data.csv", "text/csv")
        st.caption("Settings changed after building? Click the build button again to refresh the files.")
