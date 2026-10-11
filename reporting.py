"""Figures (matplotlib) and exports (Excel, Word) for KAP Survey Analyzer."""
from __future__ import annotations

import io
from datetime import date

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

import kap_core as kc  # noqa: E402

# Palette (validated: ordinal blue ramp; diverging red <-> gray <-> blue)
INK, INK2, MUTED, GRID, SURFACE = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#ffffff"
ORDINAL = {"high": "#184f95", "mid": "#3987e5", "low": "#86b6ef"}
BAR = "#2a78d6"
NEG_DARK, NEG_LIGHT, NEUTRAL, POS_LIGHT, POS_DARK = "#c4302f", "#ee8584", "#d9d8d2", "#86b6ef", "#256abf"


def _style(ax):
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#c3c2b7")
    ax.tick_params(colors=INK2, labelsize=9, length=0)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def _hex_mix(a: str, b: str, t: float) -> str:
    ca = np.array([int(a[i:i + 2], 16) for i in (1, 3, 5)])
    cb = np.array([int(b[i:i + 2], 16) for i in (1, 3, 5)])
    c = np.round(ca + (cb - ca) * t).astype(int)
    return "#" + "".join(f"{x:02x}" for x in c)


def diverging_colors(k: int) -> list[str]:
    """k ordered categories, negative -> positive."""
    half = k // 2
    neg = [_hex_mix(NEG_DARK, NEG_LIGHT, i / max(half - 1, 1)) for i in range(half)] if half > 1 else [NEG_DARK][:half]
    pos = [_hex_mix(POS_LIGHT, POS_DARK, i / max(half - 1, 1)) for i in range(half)] if half > 1 else [POS_DARK][:half]
    return neg + ([NEUTRAL] if k % 2 else []) + pos


def fig_levels(results: list[kc.DomainResult]):
    fig, ax = plt.subplots(figsize=(7.5, 0.75 * len(results) + 1.3), dpi=150)
    names = [r.cfg.name for r in results][::-1]
    for y, r in enumerate(results[::-1]):
        order = r.cfg.level_order
        cols = [ORDINAL["high"], ORDINAL["mid"], ORDINAL["low"]] if len(order) == 3 else [ORDINAL["high"], ORDINAL["low"]]
        counts = r.level.value_counts()
        left = 0.0
        for lev, col in zip(order, cols):
            w = counts.get(lev, 0) / r.n_scored * 100 if r.n_scored else 0
            ax.barh(y, w, left=left, color=col, edgecolor=SURFACE, linewidth=2, height=0.62)
            if w >= 7:
                txt = f"{lev}\n{w:.1f}%" if w >= 16 else f"{w:.0f}%"
                ax.text(left + w / 2, y, txt, ha="center", va="center", fontsize=8.5,
                        color="white" if col != ORDINAL["low"] else INK, linespacing=1.1)
            left += w
    ax.set_yticks(range(len(names)), names, fontsize=10, color=INK)
    ax.set_xlim(0, 100)
    ax.set_xlabel("% of respondents", color=INK2, fontsize=9)
    _style(ax)
    ax.set_title("KAP levels by domain", loc="left", fontsize=11, color=INK, pad=10)
    handles = [plt.Rectangle((0, 0), 1, 1, color=ORDINAL[k]) for k in ("high", "mid", "low")]
    ax.legend(handles, ["Highest level", "Middle level", "Lowest level"], ncol=3, frameon=False, fontsize=8,
              loc="upper left", bbox_to_anchor=(0, -0.28 if len(results) < 3 else -0.18), labelcolor=INK2)
    fig.tight_layout()
    return fig


def fig_key_items(res: kc.DomainResult, labels: dict[str, str] | None = None):
    pct = (res.item_scores == 1).sum() / res.item_scores.notna().sum() * 100
    pct = pct.sort_values()
    fig, ax = plt.subplots(figsize=(7.5, 0.32 * len(pct) + 1.2), dpi=150)
    ax.barh(range(len(pct)), pct.values, color=BAR, height=0.6, edgecolor=SURFACE, linewidth=1)
    for i, v in enumerate(pct.values):
        ax.text(v + 1, i, f"{v:.1f}%", va="center", fontsize=8, color=INK2)
    ylab = [_short(labels.get(i, i) if labels else i) for i in pct.index]
    ax.set_yticks(range(len(pct)), ylab, fontsize=8.5, color=INK)
    ax.set_xlim(0, 108)
    ax.set_xlabel("% answering correctly", color=INK2, fontsize=9)
    _style(ax)
    ax.set_title(f"{res.cfg.name}: correct responses by item", loc="left", fontsize=11, color=INK, pad=10)
    fig.tight_layout()
    return fig


def fig_likert(df: pd.DataFrame, res: kc.DomainResult):
    cfg = res.cfg
    labs = kc._ordered_labels(df, cfg)
    nmap = {kc.norm(k): v for k, v in cfg.mapping.items() if v is not None and not pd.isna(v)}
    labs = [l for l in labs if kc.norm(l) in nmap or not np.isnan(kc._try_float(l))]
    k = len(labs)
    if k < 2:
        return None
    cols = diverging_colors(k)
    items = cfg.items
    data = []
    for it in items:
        s = df[it].map(kc.norm)
        tot = s.notna().sum()
        row = [((s == kc.norm(l)).sum() / tot * 100) if tot else 0 for l in labs]
        if it in cfg.reverse:
            row = row[::-1]  # show reverse-scored items oriented so right = favourable
        data.append(row)
    data = np.array(data)
    half = k // 2
    neg = data[:, :half].sum(1) + (data[:, half] / 2 if k % 2 else 0)
    fig, ax = plt.subplots(figsize=(8, 0.36 * len(items) + 1.6), dpi=150)
    for i in range(len(items)):
        left = -neg[i]
        for j in range(k):
            w = data[i, j]
            ax.barh(i, w, left=left, color=cols[j], edgecolor=SURFACE, linewidth=1.5, height=0.62)
            if w >= 9:
                dark = cols[j] in (NEG_DARK, POS_DARK) or j in (0, k - 1)
                ax.text(left + w / 2, i, f"{w:.0f}", ha="center", va="center", fontsize=7.5,
                        color="white" if dark else INK)
            left += w
    ax.axvline(0, color=MUTED, linewidth=0.8, zorder=0.5)
    lim = max(neg.max(), (data.sum(1) - neg).max()) + 5
    ax.set_xlim(-lim, lim)
    ticks = [t for t in range(-100, 101, 25) if abs(t) <= lim]
    ax.set_xticks(ticks, [f"{abs(t)}%" for t in ticks])
    ax.set_yticks(range(len(items)), [_short(i) + (" (R)" if i in cfg.reverse else "") for i in items],
                  fontsize=8.5, color=INK)
    ax.invert_yaxis()
    _style(ax)
    ax.set_title(f"{cfg.name}: response distribution by item", loc="left", fontsize=11, color=INK, pad=10)
    handles = [plt.Rectangle((0, 0), 1, 1, color=c) for c in cols]
    ax.legend(handles, labs, ncol=min(k, 5), frameon=False, fontsize=7.5, loc="upper center",
              bbox_to_anchor=(0.5, -0.04 - 0.6 / (0.36 * len(items) + 1.6)), labelcolor=INK2)
    if cfg.reverse:
        fig.text(0.01, 0.005, "(R) reverse-scored item, shown with its scale flipped so the right side is favourable.",
                 fontsize=7, color=MUTED)
    fig.tight_layout()
    return fig


def _short(s: str, n: int = 48) -> str:
    s = str(s)
    return s if len(s) <= n else s[: n - 1] + "…"


def fig_png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _clean(df: pd.DataFrame) -> pd.DataFrame:
    return df[[c for c in df.columns if not str(c).startswith("_")]] if df is not None else pd.DataFrame()


def build_excel(sections: list[tuple[str, pd.DataFrame]], methods: str, results: list[str],
                scored: pd.DataFrame) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        info = pd.DataFrame({"KAP Survey Analyzer": [
            f"Generated {date.today().isoformat()}", "", "Statistical analysis (methods)", methods, "",
            "Results summary", *results]})
        info.to_excel(xw, sheet_name="Summary", index=False)
        used = {"Summary"}
        for title, tab in sections:
            name = _sheet_name(title, used)
            _clean(tab).to_excel(xw, sheet_name=name, index=False)
        scored.to_excel(xw, sheet_name=_sheet_name("Scored data", used), index=False)
        head = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="184F95")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font, c.fill = head, fill
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
            ws.freeze_panes = "A2"
        s = xw.book["Summary"]
        s.column_dimensions["A"].width = 140
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    return buf.getvalue()


def _sheet_name(title: str, used: set) -> str:
    bad = '[]:*?/\\'
    base = "".join(ch for ch in title if ch not in bad)[:31]
    name, i = base, 2
    while name in used:
        name = f"{base[:28]}_{i}"
        i += 1
    used.add(name)
    return name


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def build_docx(title: str, n: int, methods: str, blocks: list[dict]) -> bytes:
    """blocks: dicts of type 'heading' | 'para' | 'table' (df, caption, note) | 'figure' (png, caption)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.9)

    doc.add_heading(title, level=0)
    p = doc.add_paragraph()
    r = p.add_run(f"KAP Survey Analyzer report · N = {n} · {date.today().strftime('%d %B %Y')}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    doc.add_heading("Statistical analysis", level=1)
    doc.add_paragraph(methods)
    doc.add_heading("Results", level=1)

    t_no = f_no = 0
    for b in blocks:
        if b["type"] == "heading":
            _orient(doc, landscape=False)
            doc.add_heading(b["text"], level=2)
        elif b["type"] == "para":
            _orient(doc, landscape=False)
            doc.add_paragraph(b["text"])
        elif b["type"] == "table":
            tab = _clean(b["df"])
            if tab.empty:
                continue
            _orient(doc, landscape=_table_width_in(tab) > 6.9)
            t_no += 1
            cap = doc.add_paragraph()
            cr = cap.add_run(f"Table {t_no}. ")
            cr.bold = True
            cap.add_run(b["caption"])
            cap.paragraph_format.space_after = Pt(4)
            cap.paragraph_format.keep_with_next = True
            _three_line_table(doc, tab, Pt)
            if b.get("note"):
                np_ = doc.add_paragraph()
                nr = np_.add_run(b["note"])
                nr.font.size = Pt(8)
                nr.italic = True
            doc.add_paragraph()
        elif b["type"] == "figure" and b.get("png"):
            _orient(doc, landscape=False)
            f_no += 1
            doc.add_picture(io.BytesIO(b["png"]), width=Inches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cr = cap.add_run(f"Figure {f_no}. ")
            cr.bold = True
            cap.add_run(b["caption"])
    _orient(doc, landscape=False)
    foot = doc.add_paragraph()
    fr = foot.add_run("Generated with KAP Survey Analyzer. Verify variable coding and cut-offs before publication.")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x89, 0x87, 0x81)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _table_width_in(df: pd.DataFrame) -> float:
    """Rough rendered width (inches) of a table at 8.5 pt, to decide on landscape pages."""
    total = 0.0
    for c in df.columns:
        longest_cell = max([len(str(v)) for v in df[c].tolist()] + [0])
        longest_word = max(len(w) for w in str(c).split()) if str(c).split() else 4
        total += max(longest_cell, longest_word, 4) * 0.062 + 0.15
    return total


def _orient(doc, landscape: bool):
    """Start a new section when the page orientation needs to change."""
    from docx.enum.section import WD_ORIENT, WD_SECTION

    cur = doc.sections[-1]
    is_land = cur.orientation == WD_ORIENT.LANDSCAPE
    if is_land == landscape:
        return
    new = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = new.page_width, new.page_height
    new.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if (landscape and w < h) or (not landscape and w > h):
        new.page_width, new.page_height = h, w


def _three_line_table(doc, df: pd.DataFrame, Pt):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    t = doc.add_table(rows=1, cols=len(df.columns))
    t.autofit = True

    def border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            tcPr.append(b)
        for edge, sz in edges.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
            b.append(el)

    def write(cell, text, bold=False):
        cell.text = ""
        run = cell.paragraphs[0].add_run("" if pd.isna(text) else str(text))
        run.font.size = Pt(8.5)
        run.bold = bold
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    for j, c in enumerate(df.columns):
        cell = t.rows[0].cells[j]
        write(cell, c, bold=True)
        border(cell, top=12, bottom=6)
    for i, (_, row) in enumerate(df.iterrows()):
        cells = t.add_row().cells
        for j, v in enumerate(row.values):
            if isinstance(v, float) and not pd.isna(v):
                v = (f"{v:.0f}" if float(v).is_integer() and abs(v) >= 1 else f"{v:.3f}") if abs(v) < 1e4 \
                    else f"{v:,.0f}"
            write(cells[j], v)
            if i == len(df) - 1:
                border(cells[j], bottom=12)


def test_note(tables: list[pd.DataFrame]) -> str:
    """Footnote listing only the test superscripts that appear in the tables."""
    text = " ".join(" ".join(map(str, t.astype(str).values.ravel())) for t in tables if t is not None)
    inv = {v: k for k, v in kc.TEST_MARK.items()}
    used = []
    for mark in sorted(inv):
        for tok in text.split():
            if tok.endswith(mark) and tok[:-1].replace(".", "").replace("<", "").isdigit():
                used.append(f"{mark} {inv[mark]}")
                break
    return ("Tests: " + "; ".join(used) + ".") if used else ""
